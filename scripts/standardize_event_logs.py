
from pathlib import Path
import csv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

SOURCE_FILE = (
    ROOT / "data" / "raw" / "careflow_synthetic_events 1.xlsx"
)
OUTPUT_FILE = ROOT / "data" / "processed" / "careflow_event_log.csv"

EXPECTED_COLUMNS = [
    "Event_ID",
    "Case_ID",
    "Activity",
    "Timestamp",
    "Department",
    "Doctor_ID",
    "Patient_ID",
    "Age_Group",
    "Event_Status",
    "Outcome",
    "Patient_Satisfaction_Score",
    "Inpatient_Bed_Availability",
]

OUTPUT_COLUMNS = [
    "Event_ID",
    "Case_ID",
    "Activity_Name",
    "Timestamp",
    "Department",
    "Doctor_ID",
    "Patient_ID",
    "Age_Group",
    "Event_Status",
    "Outcome",
    "Patient_Satisfaction_Score",
    "Inpatient_Bed_Availability",
]


def main():
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(f"Source not found: {SOURCE_FILE}")

    workbook = load_workbook(
        SOURCE_FILE, read_only=True, data_only=True
    )
    sheet = workbook.active

    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(iterator)]

    if headers != EXPECTED_COLUMNS:
        workbook.close()
        raise ValueError(
            "Source columns differ from the agreed schema.\n"
            f"Expected: {EXPECTED_COLUMNS}\n"
            f"Found:    {headers}"
        )

    records = []

    for row in iterator:
        if not any(value is not None for value in row):
            continue

        record = dict(zip(headers, row))

        timestamp = record["Timestamp"]
        if not hasattr(timestamp, "isoformat"):
            workbook.close()
            raise ValueError(
                f"Invalid timestamp in event {record['Event_ID']}"
            )

        record["Timestamp"] = timestamp.isoformat(
            sep=" ", timespec="seconds"
        )

        # Standardize the activity field without changing its values.
        record["Activity_Name"] = record.pop("Activity")

        records.append(record)

    workbook.close()

    # Stable chronological ordering within each case.
    records.sort(
        key=lambda r: (r["Case_ID"], r["Timestamp"], r["Event_ID"])
    )

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    with OUTPUT_FILE.open(
        "w", newline="", encoding="utf-8"
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=OUTPUT_COLUMNS,
            extrasaction="raise",
        )
        writer.writeheader()
        writer.writerows(records)

    print("STANDARDIZATION COMPLETE")
    print(f"Output: {OUTPUT_FILE}")
    print(f"Rows: {len(records)}")
    print(f"Columns: {len(OUTPUT_COLUMNS)}")
    print("Activity column renamed to Activity_Name.")
    print("Original Excel workbook remains unchanged.")


if __name__ == "__main__":
    main()
